from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

PARENT_RE = re.compile(r"^((?:DES|DPI)\d+)\s*-\s*(.+)$", re.I)
DETAIL_RE = re.compile(r"^(\d{6}[A-Z]?)\s*-\s*(.+)$", re.I)
BUD_COL_RE = re.compile(r"^BUD\s+H([12])\s+(\d{2})$", re.I)


@dataclass
class BudgetParseResult:
    parent_df: pd.DataFrame
    detail_df: pd.DataFrame
    entity: str
    scope: str
    reference_period: str
    h1_col: int
    h2_col: int


def _to_float(value):
    try:
        return float(value) if value not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0


def parse_reference_budget(path: str | Path, sheet_name: str = "IOO") -> BudgetParseResult:
    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Reference Budget sheet '{sheet_name}' was not found.")
    ws = wb[sheet_name]

    entity = str(ws["B1"].value or "").strip()
    scope = str(ws["B2"].value or "").strip()

    header_row = None
    h1_col = h2_col = None
    year = None
    for r in range(1, min(ws.max_row, 30) + 1):
        for c in range(1, min(ws.max_column, 20) + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str):
                m = BUD_COL_RE.match(v.strip().upper())
                if m:
                    header_row = r
                    year = 2000 + int(m.group(2))
                    if m.group(1) == "1":
                        h1_col = c
                    elif m.group(1) == "2":
                        h2_col = c
    if not (header_row and h1_col and h2_col and year):
        wb.close()
        raise ValueError("Could not identify Bud H1 / Bud H2 reference columns in the IOO sheet.")

    parents = []
    details = []
    current_parent = None
    for r in range(header_row + 1, ws.max_row + 1):
        label = ws.cell(r, 1).value
        if label is None:
            continue
        text = str(label).strip()
        pm = PARENT_RE.match(text)
        dm = DETAIL_RE.match(text)
        h1 = _to_float(ws.cell(r, h1_col).value)
        h2 = _to_float(ws.cell(r, h2_col).value)

        if pm:
            current_parent = pm.group(1).upper()
            # Parent rows define the management ceiling. Ignore rows with no reference values.
            if h1 != 0.0 or h2 != 0.0:
                annual = h1 + h2
                parents.append({
                    "MANAGEMENT_CODE": current_parent,
                    "DESCRIPTION": text,
                    "SOURCE_H1_KEUR": h1,
                    "SOURCE_H2_KEUR": h2,
                    "ANNUAL_SOURCE_KEUR": annual,
                    "REFERENCE_BUDGET_EUR": abs(annual) * 1000.0,
                })
            continue

        if dm and current_parent:
            account = dm.group(1).upper()
            category = dm.group(2).strip()
            annual = h1 + h2
            details.append({
                "MANAGEMENT_CODE": current_parent,
                "BUDGET_ACCOUNT": account,
                "BUDGET_CATEGORY": category,
                "SOURCE_H1_KEUR": h1,
                "SOURCE_H2_KEUR": h2,
                "ORIGINAL_BUDGET_KEUR": annual,
                "REFERENCE_BUDGET_EUR": abs(annual) * 1000.0,
            })

    wb.close()
    parent_df = pd.DataFrame(parents).drop_duplicates(subset=["MANAGEMENT_CODE"], keep="last")
    detail_df = pd.DataFrame(details)
    if parent_df.empty or detail_df.empty:
        raise ValueError("No usable Mauritius parent/detail budget rows were extracted from IOO.")

    return BudgetParseResult(
        parent_df=parent_df.reset_index(drop=True),
        detail_df=detail_df.reset_index(drop=True),
        entity=entity,
        scope=scope,
        reference_period=str(year),
        h1_col=h1_col,
        h2_col=h2_col,
    )
