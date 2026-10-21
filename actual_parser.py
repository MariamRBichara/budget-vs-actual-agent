from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from config import ACTUAL_REQUIRED_FIELDS

MONTH_RE = re.compile(r"^(JAN|FEB|MAR|APR|MAY|JUN|JUL|AUG|SEP|OCT|NOV|DEC)-(\d{2})$", re.I)
MONTH_ORDER = {m: i for i, m in enumerate("JAN FEB MAR APR MAY JUN JUL AUG SEP OCT NOV DEC".split(), start=1)}


@dataclass
class ActualParseResult:
    dataframe: pd.DataFrame
    entity: str
    actual_period: str
    month_columns: list[str]
    source_totals: dict[str, float]
    source_grand_totals: dict[str, float]
    header_row: int


def _clean_header(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _find_transaction_sheet_and_header(path: str | Path) -> tuple[str, int]:
    wb = load_workbook(path, read_only=True, data_only=True)
    required = set(ACTUAL_REQUIRED_FIELDS[:4])
    for ws in wb.worksheets:
        for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 80), values_only=True), start=1):
            values = {_clean_header(v).upper() for v in row if v is not None}
            if required.issubset(values):
                wb.close()
                return ws.title, row_idx
    wb.close()
    raise ValueError("Could not locate the Actual transaction header row.")


def _metadata(path: str | Path, sheet_name: str, header_row: int) -> dict[str, str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    meta: dict[str, str] = {}
    for r in range(1, header_row):
        key = ws.cell(r, 1).value
        val = ws.cell(r, 2).value
        if key not in (None, "") and val not in (None, ""):
            meta[str(key).strip().upper()] = str(val).strip()
    wb.close()
    return meta


def _period_label(month_columns: list[str]) -> str:
    """Return a compact chronological period label for any monthly Actual columns."""
    if not month_columns:
        return "Unknown"

    parsed = []
    for col in month_columns:
        m = MONTH_RE.match(col.upper())
        if m:
            month_num = MONTH_ORDER[m.group(1).upper()]
            year = 2000 + int(m.group(2))
            month_name = m.group(1).title()
            parsed.append((year, month_num, month_name))

    parsed.sort(key=lambda x: (x[0], x[1]))
    if not parsed:
        return "Unknown"

    start_year, _, start_month = parsed[0]
    end_year, _, end_month = parsed[-1]

    if len(parsed) == 1:
        return f"{start_month} {start_year}"
    if start_year == end_year:
        return f"{start_month}–{end_month} {start_year}"
    return f"{start_month} {start_year}–{end_month} {end_year}"


def parse_actual_excel(path: str | Path) -> ActualParseResult:
    path = str(path)
    sheet_name, header_row = _find_transaction_sheet_and_header(path)
    meta = _metadata(path, sheet_name, header_row)

    raw = pd.read_excel(path, sheet_name=sheet_name, header=header_row - 1, dtype=object)
    raw.columns = [_clean_header(c).upper() for c in raw.columns]
    raw = raw.loc[:, ~raw.columns.str.startswith("UNNAMED")]

    missing = [c for c in ACTUAL_REQUIRED_FIELDS if c not in raw.columns]
    if missing:
        raise ValueError(f"Actual file is missing required columns: {missing}")

    month_columns = [c for c in raw.columns if MONTH_RE.match(c)]
    month_columns.sort(key=lambda c: (2000 + int(MONTH_RE.match(c).group(2)), MONTH_ORDER[MONTH_RE.match(c).group(1).upper()]))
    if not month_columns:
        raise ValueError("No monthly Actual columns such as JAN-26 / FEB-26 were found.")

    # Capture a source Grand Total row if Finance supplied one, then remove non-transaction rows.
    grand_mask = raw.apply(
        lambda row: any(str(v).strip().lower() == "grand total" for v in row.tolist() if v is not None),
        axis=1,
    )
    source_grand_totals: dict[str, float] = {}
    if grand_mask.any():
        grand = raw.loc[grand_mask].iloc[-1]
        for c in month_columns:
            source_grand_totals[c] = float(pd.to_numeric(grand.get(c), errors="coerce") or 0.0)

    df = raw.loc[~grand_mask].copy()
    # A real transaction requires management code + account code. This drops trailing blanks/footers.
    df = df[df["MANAGEMENT_CODE"].notna() & df["ACCOUNT_CODE"].notna()].copy()

    # Preserve the original Excel row number. Data row 0 corresponds to header_row + 1.
    df.insert(0, "SOURCE_ROW", [header_row + 1 + i for i in df.index])

    for c in month_columns:
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0.0).astype(float)

    for c in ["MANAGEMENT_CODE", "LIB4", "ACCOUNT_CODE", "DESCRIPTION", "PARTNER_CODE", "PO_NUMBER"]:
        df[c] = df[c].fillna("").astype(str).str.strip()

    df["LIB4"] = df["LIB4"].replace("", "(blank)")
    df["ACCOUNT_CODE"] = df["ACCOUNT_CODE"].str.replace(r"\.0$", "", regex=True)
    df["ACTUAL_TOTAL"] = df[month_columns].sum(axis=1)

    source_totals = {c: float(df[c].sum()) for c in month_columns}
    source_totals["TOTAL"] = float(df["ACTUAL_TOTAL"].sum())

    entity = meta.get("ENTITY_CTRY_DESC") or meta.get("ENTITY_DESC") or "Unknown"
    return ActualParseResult(
        dataframe=df.reset_index(drop=True),
        entity=entity,
        actual_period=_period_label(month_columns),
        month_columns=month_columns,
        source_totals=source_totals,
        source_grand_totals=source_grand_totals,
        header_row=header_row,
    )
